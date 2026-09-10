# -*- coding: utf-8 -*-
# ============================================================
#  Docker Notes -> PPTX  (content: exactly from your DOCX)
#  Run:  pip install python-pptx
#        python docker_pptx.py
# ============================================================
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ---------- colors ----------
NAVY=RGBColor(0x0B,0x21,0x4A); BLUE=RGBColor(0x1D,0x63,0xED)
MBLUE=RGBColor(0x17,0x4E,0xA6); SKY=RGBColor(0x5B,0xA8,0xF5)
LBLUE=RGBColor(0xE8,0xF1,0xFF); TEAL=RGBColor(0x0E,0x7C,0x86)
ORANGE=RGBColor(0xF5,0xA6,0x23); LORNG=RGBColor(0xFD,0xEB,0xD0)
LORNG2=RGBColor(0xF8,0xCB,0xAD); DORNG=RGBColor(0xC5,0x5A,0x11)
GREEN=RGBColor(0x18,0xA0,0x5E); GGREEN=RGBColor(0x70,0xAD,0x47)
RED=RGBColor(0xD9,0x30,0x25); DRED=RGBColor(0x8B,0x1E,0x1E)
GOLD=RGBColor(0xFF,0xC0,0x3D); DARK=RGBColor(0x21,0x2B,0x36)
GREY=RGBColor(0xF2,0xF4,0xF8); WHITE=RGBColor(0xFF,0xFF,0xFF)
MONO="Consolas"

prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
BLANK=prs.slide_layouts[6]

# ---------- helpers ----------
def slide(bg=WHITE):
    s=prs.slides.add_slide(BLANK)
    r=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,prs.slide_width,prs.slide_height)
    r.fill.solid(); r.fill.fore_color.rgb=bg; r.line.fill.background(); r.shadow.inherit=False
    return s

def box(s,x,y,w,h,text="",fill=BLUE,line=None,round_=True,size=14,color=WHITE,
        bold=True,align=PP_ALIGN.CENTER,font="Calibri",shape=None):
    if shape is None:
        shape=MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE
    shp=s.shapes.add_shape(shape,Inches(x),Inches(y),Inches(w),Inches(h))
    shp.shadow.inherit=False
    if fill is None: shp.fill.background()
    else: shp.fill.solid(); shp.fill.fore_color.rgb=fill
    if line is None: shp.line.fill.background()
    else: shp.line.color.rgb=line; shp.line.width=Pt(1.25)
    tf=shp.text_frame; tf.word_wrap=True
    tf.margin_left=Pt(6); tf.margin_right=Pt(6); tf.margin_top=Pt(3); tf.margin_bottom=Pt(3)
    tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    if text:
        p=tf.paragraphs[0]; p.alignment=align; r=p.add_run(); r.text=text
        r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color; r.font.name=font
    return shp

def par(shp,text,size=14,color=WHITE,bold=False,italic=False,align=PP_ALIGN.CENTER,
        font="Calibri",space_before=3):
    p=shp.text_frame.add_paragraph(); p.alignment=align; p.space_before=Pt(space_before)
    r=p.add_run(); r.text=text
    r.font.size=Pt(size); r.font.bold=bold; r.font.italic=italic
    r.font.color.rgb=color; r.font.name=font
    return p

def tbox(s,x,y,w,h,lines,size=16,color=DARK,bold=False,align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP,font="Calibri"):
    tb=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=tb.text_frame; tf.word_wrap=True; tf.vertical_anchor=anchor
    tf.margin_left=Pt(4); tf.margin_right=Pt(4); tf.margin_top=Pt(2); tf.margin_bottom=Pt(2)
    first=True
    for item in lines:
        text,ov=(item if isinstance(item,tuple) else (item,{}))
        p=tf.paragraphs[0] if first else tf.add_paragraph(); first=False
        p.alignment=ov.get("align",align); p.space_after=Pt(ov.get("space_after",5))
        if "space_before" in ov: p.space_before=Pt(ov["space_before"])
        r=p.add_run(); r.text=text; f=r.font
        f.size=Pt(ov.get("size",size)); f.bold=ov.get("bold",bold)
        f.italic=ov.get("italic",False); f.color.rgb=ov.get("color",color); f.name=ov.get("font",font)
    return tb

def title(s,text,color=RED,size=27):
    tbox(s,0.55,0.22,12.25,0.8,[(text,{"size":size,"bold":True,"color":color})])

def arrow(s,x,y,w=0.5,h=0.32,fill=NAVY,shape=MSO_SHAPE.RIGHT_ARROW):
    a=s.shapes.add_shape(shape,Inches(x),Inches(y),Inches(w),Inches(h))
    a.fill.solid(); a.fill.fore_color.rgb=fill; a.line.fill.background(); a.shadow.inherit=False
    return a

def label(s,x,y,w,text,size=12,color=DARK,bold=False,align=PP_ALIGN.CENTER,italic=False):
    return tbox(s,x,y,w,0.35,[(text,{"size":size,"color":color,"bold":bold,
             "align":align,"italic":italic})])

def chev(s,x,y,w,h,text,fill,size=15):
    shp=s.shapes.add_shape(MSO_SHAPE.CHEVRON,Inches(x),Inches(y),Inches(w),Inches(h))
    shp.fill.solid(); shp.fill.fore_color.rgb=fill; shp.line.fill.background()
    shp.shadow.inherit=False
    tf=shp.text_frame; tf.word_wrap=True; tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    p=tf.paragraphs[0]; p.alignment=PP_ALIGN.CENTER; r=p.add_run(); r.text=text
    r.font.size=Pt(size); r.font.bold=True; r.font.color.rgb=WHITE; r.font.name="Calibri"

def qa_box(s,x,y,w,h,question,answer):
    b=box(s,x,y,w,h,"",fill=LBLUE,line=MBLUE)
    p=b.text_frame.paragraphs[0]; p.alignment=PP_ALIGN.LEFT
    r=p.add_run(); r.text=question
    r.font.size=Pt(16); r.font.bold=True; r.font.color.rgb=NAVY; r.font.name="Calibri"
    par(b,answer,size=15,color=DARK,italic=True,align=PP_ALIGN.LEFT,space_before=6)
    return b

# ============ SLIDE 1 : Title ============
s=slide(NAVY)
box(s,0,3.02,13.333,0.09,"",fill=BLUE,round_=False)
tbox(s,0.5,1.55,12.33,1.4,[("Docker",{"size":76,"bold":True,"color":WHITE,"align":PP_ALIGN.CENTER})])
tbox(s,0.5,3.35,12.33,0.7,[("What is Docker  •  Why Docker  •  How Docker works  •  Docker vs VM  •  Containers  •  Images  •  Lifecycle  •  Isolation  •  Benefits & Limitations  •  Virtualization",{"size":14,"color":SKY,"align":PP_ALIGN.CENTER})])
for i in range(4): box(s,5.49+i*0.62,4.6,0.5,0.5,"",fill=BLUE)
box(s,5.49,5.18,2.36,0.18,"",fill=ORANGE,round_=False)

# ============ SLIDE 2 : 1. What is Docker ============
s=slide(); title(s,"1 — What is Docker?")
tbox(s,0.6,1.15,12.15,1.5,[("Docker is an open-source containerization platform that packages an application along with its dependencies, libraries, and configuration into a lightweight container. This allows the application to run consistently across different environments and solves the 'it works on my machine' problem.",{"size":18})])
items=[("Docker",NAVY),("Package App + Dependencies",BLUE),("Container",ORANGE),("Consistent Environment",GREEN)]
bw=(12.15-3*0.5)/4
for i,(t,c) in enumerate(items):
    x=0.6+i*(bw+0.5); box(s,x,3.4,bw,1.0,t,fill=c,size=15)
    if i<3: arrow(s,x+bw+0.08,3.76,0.34,0.28)

# ============ SLIDE 3 : 2. Why do we use Docker ============
s=slide(); title(s,"2 — Why do we use Docker?")
tbox(s,0.6,1.15,12.15,1.6,[("We use Docker to ensure consistency across different environments. It packages an application along with its dependencies and configurations into a lightweight container, so it runs consistently in development, testing, and production. It also provides isolation, portability, and faster deployment.",{"size":18})])
for i,t in enumerate(["Consistency","Isolation","Lightweight","Portability"]):
    chev(s,0.6+i*3.02,3.15,3.35,0.85,t,[MBLUE,BLUE,TEAL,GREEN][i])
qa_box(s,0.6,4.55,12.15,2.2,
 'If interviewer says: "Explain a little more"  →  I will answer:',
 '"Docker solves the \'it works on my machine\' problem because the same container can run across different environments without dependency or configuration conflicts."')

# ============ SLIDE 4 : 3. How does Docker work ============
s=slide(); title(s,"3 — How does Docker work?")
tbox(s,0.6,1.0,12.15,1.6,[("Docker works by using a containerization approach. First, we define the application and its dependencies in a Dockerfile. Docker uses the Dockerfile to build a Docker Image, which is a read-only blueprint of the application. When we run the image, Docker creates a Docker Container, which is the running instance of that application. The Docker Client communicates with the Docker Daemon, which manages the containers.",{"size":16})])
b1=box(s,0.6,2.7,3.5,1.3,"Dockerfile",fill=NAVY,size=20); par(b1,"Instructions",size=14,color=SKY,bold=True)
b2=box(s,4.9,2.7,3.5,1.3,"Docker Image",fill=BLUE,size=20); par(b2,"Blueprint",size=14,color=LBLUE,bold=True)
b3=box(s,9.2,2.7,3.5,1.3,"Docker Container",fill=ORANGE,size=20); par(b3,"Running Application",size=14,bold=True)
arrow(s,4.2,3.2,0.6,0.3); arrow(s,8.5,3.2,0.6,0.3)
c1=box(s,2.2,4.2,2.9,0.55,"Docker Client",fill=SKY,color=NAVY,size=13)
arrow(s,5.2,4.33,0.4,0.28)
c2=box(s,5.7,4.2,2.9,0.55,"Docker Daemon",fill=MBLUE,size=13)
arrow(s,8.7,4.33,0.4,0.28)
c3=box(s,9.2,4.2,2.9,0.55,"Manages Containers",fill=NAVY,size=13)
qa_box(s,0.6,5.05,12.15,1.95,
 'If interviewer asks: "Why is Docker lightweight?"  →  I will answer:',
 '"Unlike a traditional VM, Docker containers share the host OS kernel instead of running a complete guest operating system. That\'s why containers are lightweight and start much faster."')

# ============ SLIDE 5 : 4. Docker vs VM ============
s=slide(); title(s,"4 — Docker vs Virtual Machines? Which is a Better Choice? and why?",size=24)
tbox(s,0.6,1.05,12.15,1.55,[("Neither is universally better; it depends on the use case. For most modern application development and deployment, I would prefer Docker because containers are lightweight, start quickly, use fewer resources, and provide good portability. However, if I need a complete operating system, stronger isolation, or to run different operating systems on the same physical machine, I would choose a Virtual Machine.",{"size":16,"bold":True})])
box(s,0.6,2.7,5.9,0.55,"Docker  →  Applications",fill=BLUE,size=16)
cA=box(s,0.6,3.25,5.9,1.95,"•  Lightweight",fill=LBLUE,line=BLUE,size=15,color=DARK,bold=False,align=PP_ALIGN.LEFT)
for t in ["•  Fast startup","•  Less resource usage","•  Easy deployment & scaling"]:
    par(cA,t,size=15,color=DARK,align=PP_ALIGN.LEFT,space_before=5)
box(s,6.85,2.7,5.9,0.55,"VM  →  Complete Operating Systems",fill=ORANGE,size=16)
cB=box(s,6.85,3.25,5.9,1.95,"•  Stronger isolation",fill=LORNG,line=ORANGE,size=15,color=DARK,bold=False,align=PP_ALIGN.LEFT)
for t in ["•  Full OS environment","•  More resource-intensive","•  Useful for different OS requirements"]:
    par(cB,t,size=15,color=DARK,align=PP_ALIGN.LEFT,space_before=5)
box(s,0.6,5.55,12.15,1.15,'Golden line:   "Docker is better for application portability and efficiency; VMs are better when you need full OS-level isolation and flexibility."',fill=GOLD,color=NAVY,size=16)

# ============ SLIDE 6 : 4. Interview answer ============
s=slide(); title(s,"4 — Docker vs VM  →  Interview Answer")
qa_box(s,0.6,1.1,12.15,1.9,'If interviewer asks: "So, which one would you choose?"  →  I will answer:',
 '"For deploying a typical web application or microservice, I would choose Docker. For running a completely separate operating system or when stronger isolation is required, I would choose a VM."')
box(s,0.6,3.5,4.6,0.9,"Typical web application / microservice",fill=LBLUE,line=BLUE,color=NAVY,size=14)
arrow(s,5.35,3.79,0.5,0.3); box(s,6.0,3.5,3.4,0.9,"Docker",fill=BLUE,size=18)
box(s,0.6,4.9,4.6,0.9,"Separate OS / stronger isolation needed",fill=LORNG,line=ORANGE,color=DARK,size=14)
arrow(s,5.35,5.19,0.5,0.3); box(s,6.0,4.9,3.4,0.9,"Virtual Machine",fill=ORANGE,size=18)

# ============ SLIDE 7 : VM vs Docker FIGURE ============
s=slide(); title(s,"Figure — Virtual Machine vs Docker Containerization",size=24)
label(s,0.6,1.0,5.9,"Virtual Machine",size=16,bold=True)
label(s,6.85,1.0,5.9,"Docker Containerization",size=16,bold=True)
for i in range(2):
    x=1.1+i*2.6
    box(s,x,1.45,2.2,0.42,"VM %d"%(i+1),fill=LORNG2,color=DARK,size=12,round_=False)
    box(s,x,1.87,2.2,0.42,"APP %d"%(i+1),fill=LORNG,color=DARK,size=12,round_=False)
    box(s,x,2.29,2.2,0.42,"Bin/Lib",fill=LORNG,color=DARK,size=12,round_=False)
    box(s,x,2.71,2.2,0.78,"Guest Operating System",fill=DORNG,size=12,round_=False)
box(s,0.85,3.75,5.45,0.5,"Hypervisor",fill=SKY,color=NAVY,size=13,round_=False)
box(s,0.85,4.35,5.45,0.5,"Infrastructure (Physical Server)",fill=GGREEN,size=13,round_=False)
for i in range(4):
    x=6.95+i*1.47
    box(s,x,1.45,1.32,0.4,"Container %d"%(i+1),fill=LORNG2,color=DARK,size=9,round_=False)
    box(s,x,1.85,1.32,0.45,"APP %d"%(i+1),fill=LORNG,color=DARK,size=11,round_=False)
    box(s,x,2.30,1.32,0.45,"Bin/Lib",fill=LORNG,color=DARK,size=11,round_=False)
box(s,6.85,3.0,5.9,0.5,"Docker Engine",fill=DORNG,size=13,round_=False)
box(s,6.85,3.6,5.9,0.5,"Host Operating System",fill=SKY,color=NAVY,size=13,round_=False)
box(s,6.85,4.2,5.9,0.5,"Infrastructure (Physical Server)",fill=GGREEN,size=13,round_=False)
tbox(s,0.6,5.15,5.9,1.6,[("A VM virtualizes the hardware itself — every VM carries its own complete, independent operating system (more resource-intensive).",{"size":13,"align":PP_ALIGN.CENTER})])
tbox(s,6.85,5.15,5.9,1.6,[("Containers share the host machine's OS kernel through the Docker Engine — lightweight, fast to start.",{"size":13,"align":PP_ALIGN.CENTER})])

# ============ SLIDE 8 : 5. What is a Container ============
s=slide(); title(s,"5 — What is a Container? How Container works?")
tbox(s,0.6,1.1,12.15,1.15,[("A container is a lightweight and isolated environment that packages an application with all its required dependencies, so it can run consistently across different environments.",{"size":18,"bold":True})])
box(s,2.6,2.5,8.1,0.8,"Container  =  Application  +  Dependencies  +  Isolation",fill=NAVY,size=19)
tbox(s,0.6,3.6,12.15,0.5,[("Docker — A Docker container is a running instance of a Docker image.",{"size":16,"align":PP_ALIGN.CENTER,"color":MBLUE,"bold":True})])
label(s,0.6,4.35,12.15,"Easy example",size=15,bold=True,color=TEAL)
box(s,2.35,4.8,3.6,0.9,"Image  =  Blueprint",fill=BLUE,size=17)
arrow(s,6.1,5.09,0.55,0.3)
box(s,6.8,4.8,4.2,0.9,"Container  =  Running Application",fill=ORANGE,size=17)

# ============ SLIDE 9 : Container stack + Shipping Analogy ============
s=slide(); title(s,"Container on the machine  +  The Shipping Analogy",size=24)
label(s,1.0,0.98,4.4,"Container",size=15,bold=True,color=MBLUE)
box(s,1.0,1.35,4.4,1.55,"",fill=LORNG,line=DORNG)
box(s,1.4,1.55,3.6,0.45,"APP",fill=LORNG2,color=DARK,size=12,round_=False)
box(s,1.4,2.05,3.6,0.45,"Bin/Lib",fill=LORNG2,color=DARK,size=12,round_=False)
box(s,1.0,3.05,4.4,0.5,"Docker Engine",fill=DORNG,size=13,round_=False)
box(s,1.0,3.65,4.4,0.5,"Host Operating System",fill=SKY,color=NAVY,size=13,round_=False)
box(s,1.0,4.25,4.4,0.5,"Hardware",fill=GGREEN,size=13,round_=False)
a=box(s,6.2,1.35,6.5,3.4,"The Shipping Analogy",fill=NAVY,size=20)
par(a,"A shipping container holds goods in a standard box. Any ship, truck, or crane can move it — no one cares what's inside. Software containers work the same way: a standard box for code that any server can run.",size=15,space_before=10)
g=box(s,0.6,5.15,12.15,1.6,"",fill=LBLUE,line=MBLUE)
p=g.text_frame.paragraphs[0]; p.alignment=PP_ALIGN.LEFT
r=p.add_run(); r.text="Short answer:  A container is a lightweight, isolated environment used to run an application with all its dependencies."
r.font.size=Pt(16); r.font.bold=True; r.font.color.rgb=NAVY; r.font.name="Calibri"
par(g,'"Build once, run anywhere."  —  The defining promise of container technology.',size=15,color=DARK,align=PP_ALIGN.LEFT,space_before=6)

# ============ SLIDE 10 : Container works ============
s=slide(); title(s,"Container works")
tbox(s,0.6,1.05,12.15,0.9,[("Containers don't virtualize hardware — they share the host machine's OS kernel while staying isolated from each other. This makes them fast to start and light on resources.",{"size":16})])
label(s,0.6,2.0,12.15,"Isolated Containers",size=13,bold=True,italic=True)
for i,name in enumerate(["App A","App B","App C"]):
    x=3.57+i*2.2; box(s,x,2.45,1.9,0.75,name,fill=ORANGE,size=15)
    arrow(s,x+0.8,3.3,0.3,0.35,shape=MSO_SHAPE.DOWN_ARROW)
box(s,2.4,3.8,8.5,0.6,"Container Engine  (e.g. Docker)",fill=TEAL,size=15,round_=False)
box(s,2.4,4.55,8.5,0.6,"Host Operating System (Kernel)",fill=MBLUE,size=15,round_=False)
box(s,2.4,5.3,8.5,0.6,"Infrastructure / Hardware",fill=NAVY,size=15,round_=False)

# ============ SLIDE 11 : Containers vs Application ============
s=slide(); title(s,"Containers vs. Application")
tbox(s,0.6,1.05,12.15,0.6,[("These terms get mixed up constantly. An application is the software itself; a container is the packaging that runs it.",{"size":16})])
ap=box(s,0.9,1.95,5.6,4.6,"Application",fill=WHITE,line=MBLUE,size=19,color=NAVY)
for t in ["•  The actual program or code you write (e.g. a Django website)",
          "•  Defines business logic and features",
          "•  Needs a runtime, libraries, and OS to function",
          "•  Can run directly on a machine — or inside a container"]:
    par(ap,t,size=14.5,color=DARK,align=PP_ALIGN.LEFT,space_before=8)
ct=box(s,6.85,1.95,5.6,4.6,"Container",fill=NAVY,size=19)
for t in ["•  The isolated environment that runs the application",
          "•  Bundles the app with its dependencies & config",
          "•  Ensures consistent behavior across machines",
          "•  One container can run one or more processes"]:
    par(ct,t,size=14.5,align=PP_ALIGN.LEFT,space_before=8)

# ============ SLIDE 12 : 6. Docker Image ============
s=slide(); title(s,"6 — What is a Docker Image?")
tbox(s,0.6,1.1,12.15,1.25,[("A Docker Image is a read-only template that contains an application's code, runtime, libraries, dependencies, and configuration needed to run the application. A Docker Container is created from this image.",{"size":18,"bold":True})])
b1=box(s,0.6,2.75,3.7,1.4,"Dockerfile",fill=NAVY,size=19); par(b1,"text file with instructions",size=13,color=SKY,bold=True)
b2=box(s,4.8,2.75,3.7,1.4,"Docker Image",fill=BLUE,size=19); par(b2,"application's blueprint / package",size=13,color=LBLUE,bold=True)
b3=box(s,9.0,2.75,3.7,1.4,"Docker Container",fill=ORANGE,size=19); par(b3,"image's running instance",size=13,bold=True)
arrow(s,4.38,3.3,0.35,0.28); arrow(s,8.58,3.3,0.35,0.28)
box(s,1.8,4.9,9.7,1.0,"One-line answer:   A Docker Image is a read-only blueprint used to create Docker Containers.",fill=LBLUE,line=BLUE,color=NAVY,size=17)

# ============ SLIDE 13 : 7. Lifecycle ============
s=slide(); title(s,"7 — Docker container lifecycle")
tbox(s,0.6,1.05,12.15,0.5,[("Every container moves through a predictable set of states from creation to removal.",{"size":16})])
states=[("Created","Container is set up but not running",MBLUE),
        ("Running","Process is active and executing",GREEN),
        ("Paused","Execution frozen, state retained",ORANGE),
        ("Stopped","Process ended, filesystem remains",RGBColor(0x5A,0x63,0x6B)),
        ("Removed","Container deleted from the system",RED)]
w=2.25
for i,(n_,d_,c_) in enumerate(states):
    x=0.6+i*(w+0.22)
    c=box(s,x,1.9,w,1.9,n_,fill=c_,size=18); par(c,d_,size=12.5,space_before=4)
    if i<4: arrow(s,x+w+0.02,2.7,0.18,0.28)
box(s,1.3,4.35,10.7,0.75,"docker create   →   docker start   →   docker pause / stop   →   docker rm",
    fill=GREY,color=NAVY,size=16,bold=False,font=MONO)

# ============ SLIDE 14 : State transition diagram ============
s=slide(); title(s,"Lifecycle — State Transition Diagram")
box(s,0.9,2.8,2.4,1.0,"Created",fill=MBLUE,size=18)
box(s,5.45,2.8,2.4,1.0,"Running",fill=GREEN,size=18)
box(s,10.0,2.8,2.4,1.0,"Paused",fill=ORANGE,size=18)
box(s,5.45,5.15,2.4,1.0,"Stopped",fill=DRED,size=18)
arrow(s,3.5,3.16,1.75,0.28); label(s,3.4,2.72,2.0,"docker start",size=12,bold=True,color=NAVY)
arrow(s,8.05,2.75,1.75,0.26); label(s,8.0,2.32,1.9,"docker pause",size=12,bold=True,color=NAVY)
arrow(s,8.05,3.6,1.75,0.26,shape=MSO_SHAPE.LEFT_ARROW); label(s,7.95,3.95,2.0,"docker unpause",size=12,bold=True,color=NAVY)
arrow(s,6.15,4.0,0.3,0.95,shape=MSO_SHAPE.DOWN_ARROW); label(s,4.35,4.28,1.7,"docker stop",size=12,bold=True,color=NAVY,align=PP_ALIGN.RIGHT)
arrow(s,6.85,4.0,0.3,0.95,shape=MSO_SHAPE.UP_ARROW); label(s,7.3,4.28,1.7,"docker start",size=12,bold=True,color=NAVY,align=PP_ALIGN.LEFT)

# ============ SLIDE 15 : 8. Container isolation ============
s=slide(); title(s,"8 — Container isolation?")
tbox(s,0.6,1.05,12.15,0.65,[("Isolation is what makes containers safe to run side-by-side. Each container gets its own view of the system, powered by Linux kernel features.",{"size":16})])
cards=[("Namespaces","Isolate what a container can see — its own processes, network, hostname, and file mounts.",MBLUE),
       ("Control Groups (cgroups)","Limit and meter how much CPU, memory, and I/O a container is allowed to consume.",TEAL),
       ("Filesystem Isolation","Each container has its own writable layer — changes never leak into other containers.",NAVY)]
for i,(h_,d_,c_) in enumerate(cards):
    x=0.6+i*(3.85+0.3)
    box(s,x,2.0,3.85,0.65,h_,fill=c_,size=16)
    box(s,x,2.65,3.85,2.2,d_,fill=WHITE,line=c_,size=14.5,color=DARK,bold=False,align=PP_ALIGN.LEFT)

# ============ SLIDE 16 : 9. Benefits ============
s=slide(); title(s,"9 — Benefits of containers?")
bens=[("Fast Startup","Boot in seconds, not minutes — no full OS to load.",GREEN),
      ("Lightweight","Share the host kernel, using far less resources than a VM.",BLUE),
      ("Portable","Runs identically on a laptop, server, or the cloud.",TEAL),
      ("Consistent","\"Works on my machine\" problems disappear.",MBLUE),
      ("Scalable","Spin up dozens of identical containers on demand.",NAVY),
      ("Isolated","Apps don't interfere with each other on the host.",ORANGE)]
for i,(t_,d_,c_) in enumerate(bens):
    x=0.6+(i%3)*(3.85+0.3); y=1.35+(i//3)*(2.15+0.35)
    box(s,x,y,3.85,0.55,t_,fill=c_,size=15)
    box(s,x,y+0.55,3.85,1.6,d_,fill=WHITE,line=c_,size=13.5,color=DARK,bold=False,align=PP_ALIGN.LEFT)

# ============ SLIDE 17 : 10. Limitations ============
s=slide(); title(s,"10 — Limitations of Containers?")
tbox(s,0.6,1.0,12.15,0.6,[("Containers solve many problems — but they are not a universal answer. Understanding the limits matters just as much as the benefits.",{"size":15.5})])
lims=[("Shared Kernel Risk","All containers on a host share one OS kernel — a kernel-level exploit can affect every container."),
      ("Not Full Isolation","Weaker security boundary than a virtual machine; not ideal for running fully untrusted code."),
      ("Persistent Data Needs Care","Containers are ephemeral by default — data must be deliberately stored in volumes."),
      ("Windows/Linux Limits","A Linux container needs a Linux kernel; cross-OS containers require extra tooling.")]
for i,(t_,d_) in enumerate(lims):
    x=0.6+(i%2)*(5.9+0.35); y=1.85+(i//2)*(1.9+0.3)
    box(s,x,y,5.9,0.55,t_,fill=RED,size=15)
    box(s,x,y+0.55,5.9,1.35,d_,fill=WHITE,line=RED,size=13.5,color=DARK,bold=False,align=PP_ALIGN.LEFT)

# ============ SLIDE 18 : 11. Key takeaways ============
s=slide(); title(s,"11 — Key takeaways")
takes=["A container packages an app with everything it needs to run.",
       "Containers share the host OS kernel — lighter and faster than VMs.",
       "An application is the code; a container is how it's packaged & run.",
       "Every container moves through: created → running → stopped → removed.",
       "Namespaces & cgroups give containers strong (not perfect) isolation.",
       "Fast, portable, and consistent — but not a full security sandbox."]
tcols=[MBLUE,BLUE,ORANGE,NAVY,TEAL,GREEN]
for i,t_ in enumerate(takes):
    x=0.6+(i%2)*(5.9+0.35); y=1.3+(i//2)*(1.5+0.35)
    box(s,x,y+0.35,0.8,0.8,str(i+1),fill=tcols[i],size=20,shape=MSO_SHAPE.OVAL)
    box(s,x+0.8,y,5.1,1.5,t_,fill=WHITE,line=tcols[i],size=14.5,color=DARK,bold=False,align=PP_ALIGN.LEFT)

# ============ SLIDE 19 : What Is Virtualization ============
s=slide(); title(s,"What Is Virtualization?")
tbox(s,0.6,1.05,12.15,1.0,[("Virtualization is the technology that lets one physical machine act like many. A software layer called a hypervisor divides real hardware — CPU, memory, storage — into multiple independent, isolated environments.",{"size":17})])
tbox(s,0.6,2.1,12.15,0.85,[("One machine, many worlds.",{"size":19,"bold":True,"color":NAVY}),
                            ("The core idea behind every VM, cloud instance, and container.",{"size":13.5})])
for i,t_ in enumerate(["One Physical Machine","Hypervisor","Multiple VMs"]):
    chev(s,0.9+i*4.0,3.15,4.35,0.8,t_,[NAVY,BLUE,TEAL][i])
tbox(s,0.9,4.25,11.5,1.6,[
 ("•  Physical Machine  =  real hardware",{"size":15,"bold":True}),
 ("•  Hypervisor  =  divides hardware resources",{"size":15,"bold":True}),
 ("•  VM  =  independent virtual computer",{"size":15,"bold":True}),
 ("•  Each VM has its own Guest OS",{"size":15,"bold":True})])
box(s,1.6,6.05,10.1,1.0,'One-line version:  "Virtualization allows one physical machine to run multiple isolated virtual machines using a hypervisor."',fill=LBLUE,line=BLUE,color=NAVY,size=15.5)

# ============ SLIDE 20 : Virtualization figure ============
s=slide(); title(s,"Virtualization — One Machine, Many Worlds")
label(s,4.67,1.05,4.0,"Physical Machine",size=15,bold=True,color=NAVY)
box(s,4.67,1.5,4.0,0.8,"CPU  ·  Memory  ·  Storage",fill=NAVY,size=16)
label(s,4.67,2.45,4.0,"▼  divides into  ▼",size=13,italic=True)
for i,(n_,c_) in enumerate([("Env A",MBLUE),("Env B",TEAL),("Env C",ORANGE)]):
    box(s,1.9+i*3.35,3.1,3.0,1.3,n_,fill=c_,size=18)
tbox(s,0.6,4.8,12.15,0.6,[("Each environment believes it has the whole machine to itself.",{"size":15,"align":PP_ALIGN.CENTER,"italic":True})])

# ============ SLIDE 21 : VM Architecture ============
s=slide(); title(s,"8 — Virtual Machine Architecture")
tbox(s,0.6,1.0,12.15,0.6,[("A VM virtualizes the hardware itself — so every VM carries its own complete, independent operating system.",{"size":16,"bold":True})])
label(s,0.6,1.7,12.15,"Each VM ships a FULL guest operating system",size=12.5,italic=True)
for i in range(3):
    v=box(s,1.2+i*3.9,2.15,3.3,1.05,"VM %d"%(i+1),fill=ORANGE,size=16)
    par(v,"Guest OS + App",size=12.5,bold=False)
box(s,0.9,3.5,11.5,0.55,"Hypervisor  (e.g. VMware, VirtualBox)",fill=TEAL,size=14,round_=False)
box(s,0.9,4.15,11.5,0.55,"Host Operating System",fill=MBLUE,size=14,round_=False)
box(s,0.9,4.8,11.5,0.55,"Physical Hardware",fill=NAVY,size=14,round_=False)

# ============ SLIDE 22 : Containers skip guest OS ============
s=slide(); title(s,"Containers skip the Guest OS")
tbox(s,0.6,1.0,12.15,0.6,[("Containers skip the guest OS entirely — they share the host kernel directly through a container engine.",{"size":16,"bold":True})])
label(s,0.6,1.7,12.15,"No guest OS needed — containers share the SAME kernel, so more apps fit in the same space",size=12.5,italic=True)
for i,n_ in enumerate(["App A","App B","App C","App D"]):
    box(s,1.35+i*2.75,2.15,2.45,0.85,n_,fill=ORANGE,size=15)
box(s,0.9,3.3,11.5,0.55,"Container Engine  (e.g. Docker)",fill=TEAL,size=14,round_=False)
box(s,0.9,3.95,11.5,0.55,"Host Operating System (Kernel)",fill=MBLUE,size=14,round_=False)
box(s,0.9,4.6,11.5,0.55,"Physical Hardware",fill=NAVY,size=14,round_=False)

# ============ SLIDE 23 : VM vs Container ============
s=slide(); title(s,"Virtual Machine  vs  Container")
vm=box(s,0.9,1.3,5.6,5.3,"Virtual Machine",fill=NAVY,size=20)
for t_ in ["•  Virtualizes hardware","•  Full guest OS per instance","•  Strong isolation boundary","•  Gigabytes in size","•  Boots in ~1–2 minutes"]:
    par(vm,t_,size=16,align=PP_ALIGN.LEFT,space_before=10)
ct=box(s,6.85,1.3,5.6,5.3,"Container",fill=BLUE,size=20)
for t_ in ["•  Virtualizes the OS layer","•  Shares the host kernel","•  Lightweight isolation","•  Megabytes in size","•  Boots in ~seconds"]:
    par(ct,t_,size=16,align=PP_ALIGN.LEFT,space_before=10)

prs.save("Docker_Presentation.pptx")
print("✅ Docker_Presentation.pptx created!  (23 slides)")